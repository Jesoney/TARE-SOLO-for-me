from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 定义颜色方案
PRIMARY_COLOR = RGBColor(30, 58, 95)      # 深蓝色 #1e3a5f
ACCENT_COLOR = RGBColor(243, 156, 18)     # 橙色 #f39c12
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(51, 51, 51)
LIGHT_GRAY = RGBColor(248, 249, 250)

def add_title_slide(prs, title, subtitle):
    """添加标题页"""
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)
    
    # 背景色块
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PRIMARY_COLOR
    shape.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # 副标题
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = ACCENT_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_lines, has_subtitle=False, subtitle=""):
    """添加内容页"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # 顶部色条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_COLOR
    bar.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # 副标题（如果有）
    start_y = 1.6
    if has_subtitle and subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.333), Inches(0.4))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.font.italic = True
        start_y = 1.8
    
    # 内容
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(start_y), Inches(12.333), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, line in enumerate(content_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = line
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(12)
        p.space_after = Pt(6)
        p.level = 0 if not line.startswith("  ") else 1
    
    return slide

def add_two_column_slide(prs, title, left_title, left_content, right_title, right_content):
    """添加双栏内容页"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # 顶部色条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_COLOR
    bar.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # 左栏标题
    left_title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.8), Inches(0.5))
    tf = left_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    
    # 左栏内容
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.1), Inches(5.8), Inches(5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(left_content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(8)
    
    # 右栏标题
    right_title_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.8), Inches(0.5))
    tf = right_title_box.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    
    # 右栏内容
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(2.1), Inches(5.8), Inches(5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(right_content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(8)
    
    return slide

# ===== 开始生成PPT内容 =====

# 第1页：封面
add_title_slide(prs, 
    "个人资产管理系统", 
    "Web编程基础期末大作业 | PHP+MySQL动态网站开发")

# 第2页：目录
add_content_slide(prs, "汇报目录", [
    "1. 项目概述与主题定位",
    "2. 项目需求分析",
    "3. 功能模块设计",
    "4. 数据库设计",
    "5. AI辅助使用记录",
    "6. 系统开发流程",
    "7. 功能调试与优化",
    "8. 安全设计",
    "9. 使用心得与总结"
])

# 第3页：项目概述
add_content_slide(prs, "项目概述", [
    "系统名称：个人资产管理系统",
    "目标用户：需要管理个人资产的普通用户及系统管理员",
    "核心用途：帮助用户记录、管理和统计个人资产信息",
    "业务场景：现金、银行存款、投资、房产、车辆等各类资产的管理",
    "设计初衷：让用户清晰掌握个人财务状况，合理规划资产配置",
    "技术栈：HTML5 + CSS3 + JavaScript + PHP + MySQL",
    "界面风格：深蓝色主色调 + 橙色强调色，简洁美观、响应式设计"
])

# 第4页：需求分析 - 功能需求
add_two_column_slide(prs, "功能需求分析",
    "普通用户功能",
    [
        "用户注册与登录",
        "个人资产增删改查",
        "资产分类管理",
        "资产搜索与分页",
        "数据统计与图表展示",
        "个人中心信息修改"
    ],
    "管理员功能",
    [
        "用户管理与删除",
        "查看所有用户资产统计",
        "系统数据监控",
        "权限分级控制"
    ]
)

# 第5页：需求分析 - 非功能需求
add_content_slide(prs, "非功能需求", [
    "性能需求：页面响应速度 < 2秒，支持分页查询优化大数据量加载",
    "安全性：密码加密存储、SQL防注入、XSS防护、CSRF防护、权限控制",
    "兼容性：支持Chrome、Firefox、Edge等主流浏览器",
    "易用性：界面简洁直观，操作流程清晰，支持响应式布局",
    "可维护性：代码结构清晰，模块化设计，注释完善",
    "运行环境：PHP 7.4+、MySQL 5.7+、Apache/Nginx服务器"
])

# 第6页：功能模块图
add_content_slide(prs, "功能模块设计", [
    "用户认证模块：注册、登录、退出、个人中心",
    "资产管理模块：资产增删改查、分类筛选、搜索分页",
    "分类管理模块：资产分类的添加、删除、图标颜色自定义",
    "数据统计模块：资产总览统计、分类饼图、最近资产列表",
    "系统管理模块：用户列表、搜索、分页、删除（管理员专属）",
    "公共模块：数据库连接、权限检查、工具函数、侧边导航"
])

# 第7页：数据库设计
add_content_slide(prs, "数据库设计", [
    "用户表 (users)：id、username、password、email、phone、avatar、role、created_at",
    "  - 主键：id，索引：username、email",
    "  - 密码使用VARCHAR(255)存储加密后的哈希值",
    "",
    "资产表 (assets)：id、user_id、category_id、name、amount、description、purchase_date",
    "  - 外键：user_id → users(id)，category_id → categories(id)",
    "  - 索引：user_id、category_id、name",
    "",
    "分类表 (categories)：id、name、icon、color、created_at",
    "  - 主键：id，唯一索引：name",
    "",
    "满足第三范式，表结构合理，字段类型规范，关系正确"
])

# 第8页：AI辅助记录 - 工具与Prompt
add_content_slide(prs, "AI辅助使用记录", [
    "使用AI工具：Trae IDE内置AI助手",
    "",
    "关键Prompt示例：",
    "  - 选题阶段：'帮我设计一个PHP+MySQL的个人资产管理系统的功能模块'",
    "  - 数据库设计：'设计个人资产管理系统的数据库表结构，包含用户、资产、分类表'",
    "  - 代码开发：'帮我实现PHP的用户登录功能，包含密码加密和Session管理'",
    "  - 安全优化：'如何防止SQL注入和XSS攻击，给出PHP代码示例'",
    "  - Bug排查：'管理员登录显示密码错误，帮我排查问题'",
    "",
    "AI反馈与迭代：根据AI生成的代码进行测试、验证和修改优化"
])

# 第9页：AI辅助记录 - 流程区分
add_two_column_slide(prs, "AI辅助与个人原创区分",
    "AI辅助内容",
    [
        "代码框架和基础结构生成",
        "数据库表结构建议",
        "安全防护措施方案",
        "CSS样式布局参考",
        "Bug排查思路指导"
    ],
    "个人原创内容",
    [
        "需求分析与功能规划",
        "模块划分与业务流程设计",
        "界面配色与风格定义",
        "代码测试与逻辑验证",
        "安全策略的具体实现",
        "调试过程中的问题修复"
    ]
)

# 第10页：开发流程
add_content_slide(prs, "系统开发流程", [
    "1. 主题确定：选择个人资产管理作为系统主题",
    "2. 需求分析：明确用户角色、功能需求、非功能需求",
    "3. 模块设计：划分用户认证、资产管理、分类管理、数据统计等模块",
    "4. 数据库设计：设计users、assets、categories三张核心表",
    "5. 前端制作：首页、登录页、后台布局、响应式设计",
    "6. 后端编码：PHP业务逻辑、数据库操作、Session管理",
    "7. 前后端联调：页面跳转、数据交互、表单提交",
    "8. 功能完善：搜索、分页、统计图表、权限控制",
    "9. 安全优化：密码加密、SQL防注入、XSS防护",
    "10. 测试部署：功能测试、兼容性测试、安全测试"
])

# 第11页：调试案例
add_content_slide(prs, "功能调试案例", [
    "案例1：管理员登录密码错误",
    "  - 问题现象：使用admin/admin123登录提示密码错误",
    "  - 原因分析：数据库中插入的密码哈希值与实际密码不匹配",
    "  - 解决方法：使用PHP的password_hash()重新生成正确的密码哈希",
    "",
    "案例2：资产删除权限控制",
    "  - 问题现象：普通用户可能删除他人资产",
    "  - 原因分析：删除操作未验证资产所属用户",
    "  - 解决方法：在删除SQL中添加user_id条件限制",
    "",
    "案例3：分页参数边界处理",
    "  - 问题现象：手动输入page=0或负数导致报错",
    "  - 解决方法：使用max(1, intval())确保页码最小为1"
])

# 第12页：安全设计
add_content_slide(prs, "安全设计", [
    "密码加密：使用password_hash()和password_verify()进行密码加密验证",
    "SQL防注入：使用PDO预处理语句，所有用户输入参数化查询",
    "XSS防护：使用htmlspecialchars()对输出内容进行转义",
    "CSRF防护：生成随机Token并在表单提交时验证",
    "权限控制：Session验证登录状态，区分普通用户和管理员权限",
    "非法访问拦截：未登录用户自动跳转登录页，非管理员禁止访问管理页面",
    "表单校验：前端+后端双重校验，防空值、防非法输入、防重复提交"
])

# 第13页：系统截图展示
add_content_slide(prs, "系统页面展示", [
    "首页：系统介绍、功能特色展示、登录注册入口",
    "登录页：用户名密码输入、表单验证、错误提示",
    "后台首页：资产统计卡片、分类饼图、最近资产列表",
    "资产列表：表格展示、搜索框、分类筛选、分页导航",
    "资产添加/编辑：表单输入、分类选择、日期选择、描述文本框",
    "分类管理：分类列表、添加删除、图标颜色展示",
    "个人中心：信息修改、密码修改",
    "用户管理（管理员）：用户列表、搜索、分页、删除操作"
])

# 第14页：使用心得
add_content_slide(prs, "AI使用心得", [
    "AI在开发中的优势：",
    "  - 快速生成代码框架，提高开发效率",
    "  - 提供多种解决方案，拓宽思路",
    "  - 帮助排查错误，节省调试时间",
    "",
    "AI的局限性：",
    "  - 生成的代码可能存在逻辑漏洞",
    "  - 需要人工验证和测试",
    "  - 对复杂业务理解不够深入",
    "",
    "核心技巧：",
    "  1. 分步骤提问，从框架到细节逐步细化",
    "  2. 对AI生成的代码进行测试验证后再使用"
])

# 第15页：总结
add_content_slide(prs, "项目总结", [
    "本次作业完成了一个完整的个人资产管理系统",
    "涵盖了需求分析、数据库设计、前后端开发、安全优化全流程",
    "系统功能完整，包含用户认证、资产CRUD、分类管理、数据统计等核心功能",
    "安全措施到位，包含密码加密、SQL防注入、XSS防护、权限控制",
    "通过AI辅助提高了开发效率，同时注重个人原创内容的融入",
    "在调试过程中独立解决了密码哈希、权限控制、分页边界等问题",
    "整体达到了Web编程基础课程大作业的要求"
])

# 第16页：感谢页
add_title_slide(prs, 
    "感谢聆听", 
    "个人资产管理系统 | Web编程基础期末大作业")

# 保存PPT
output_path = "/workspace/personal-asset-manager/项目汇报PPT.pptx"
prs.save(output_path)
print(f"PPT已生成: {output_path}")
